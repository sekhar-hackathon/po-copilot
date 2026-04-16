import csv
import io
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader
from openpyxl import load_workbook


class DocumentProcessor:
    """Processes uploaded documents into text chunks for embedding."""

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def process(self, file_path: str, ext: str) -> list[str]:
        raw_text = self._extract_text(file_path, ext)
        if not raw_text.strip():
            return []
        chunks = self.splitter.split_text(raw_text)
        return chunks

    def _extract_text(self, file_path: str, ext: str) -> str:
        if ext == ".pdf":
            return self._read_pdf(file_path)
        elif ext in (".xlsx", ".xls"):
            return self._read_excel(file_path)
        elif ext == ".csv":
            return self._read_csv(file_path)
        else:
            return self._read_text(file_path)

    def _read_pdf(self, path: str) -> str:
        reader = PdfReader(path)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)

    def _read_excel(self, path: str) -> str:
        wb = load_workbook(path, read_only=True, data_only=True)
        all_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_text.append(f"--- Sheet: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                if row_text.strip(" |"):
                    all_text.append(row_text)
        wb.close()
        return "\n".join(all_text)

    def _read_csv(self, path: str) -> str:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = []
            for row in reader:
                rows.append(" | ".join(row))
        return "\n".join(rows)

    def _read_text(self, path: str) -> str:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
